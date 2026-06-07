"""
Family Meal Planner
A Streamlit app that generates weekly meal plans using the Claude API.
"""
import streamlit as st
import anthropic
from pathlib import Path
from datetime import date, timedelta
from io import BytesIO
import json
import re
import os
import base64
import requests
from dotenv import load_dotenv

load_dotenv()

# ── GitHub ──
GITHUB_REPO   = "clairechong/family-meal-planner"
HISTORY_PATH  = "meal_plan_history.md"

# ── File paths ──
APP_DIR     = Path(__file__).parent
PROJECT_DIR = APP_DIR.parent
MEMORY_DIR  = (
    Path.home()
    / ".claude"
    / "projects"
    / "X--Shared-Files-Food--Health---Wellness-1-WEEKLY-MEAL-PLANS"
    / "memory"
)


def load_text(*paths):
    """Return the contents of the first file that exists."""
    for path in paths:
        try:
            p = Path(path)
            if p.exists():
                return p.read_text(encoding="utf-8")
        except Exception:
            pass
    return ""


def get_api_key():
    key = os.getenv("ANTHROPIC_API_KEY", "")
    if not key:
        try:
            key = st.secrets["ANTHROPIC_API_KEY"]
        except Exception:
            pass
    return key


def get_github_token():
    token = os.getenv("GITHUB_TOKEN", "")
    if not token:
        try:
            token = st.secrets["GITHUB_TOKEN"]
        except Exception:
            pass
    return token


def build_history_entry(plan: list, week_start: date, week_end: date) -> str:
    if week_start.month == week_end.month:
        header = f"{week_start.strftime('%b')} {week_start.day}–{week_end.day}, {week_end.year}"
    else:
        header = f"{week_start.strftime('%b')} {week_start.day}–{week_end.strftime('%b')} {week_end.day}, {week_end.year}"
    lines = [f"### {header}"]
    for day in plan:
        day_abbr = day.get("day", "")[:3]
        dinner = day.get("dinner", "")
        if dinner:
            lines.append(f"- {day_abbr}: {dinner}")
    return "\n".join(lines)


def save_history_to_github(new_entry: str) -> tuple[bool, str]:
    token = get_github_token()
    if not token:
        return False, "No GITHUB_TOKEN found — add it to Streamlit secrets or .env."
    headers = {
        "Authorization": f"token {token}",
        "Accept": "application/vnd.github.v3+json",
    }
    url = f"https://api.github.com/repos/{GITHUB_REPO}/contents/{HISTORY_PATH}"
    resp = requests.get(url, headers=headers)
    if resp.status_code != 200:
        return False, f"Could not read history file from GitHub ({resp.status_code})."
    data = resp.json()
    current = base64.b64decode(data["content"]).decode("utf-8")
    sha = data["sha"]
    marker = "## History (most recent first)"
    if marker in current:
        before, after = current.split(marker, 1)
        updated = before + marker + "\n\n" + new_entry + "\n" + after.lstrip("\n")
    else:
        updated = new_entry + "\n\n" + current
    week_label = st.session_state.get("week_label", "new week")
    payload = {
        "message": f"Add dinner history for {week_label}",
        "content": base64.b64encode(updated.encode("utf-8")).decode("utf-8"),
        "sha": sha,
    }
    put_resp = requests.put(url, headers=headers, json=payload)
    if put_resp.status_code in (200, 201):
        st.session_state.history = updated
        return True, "History saved to GitHub ✓"
    return False, f"GitHub write failed ({put_resp.status_code})."


def build_system_prompt(notes: str, history: str) -> str:
    return f"""You are a friendly and practical meal planning assistant for a family of 3 (2 adults + 1 child) in Ontario, Canada.

FAMILY RULES AND RECIPE LIBRARY:
{notes}

RECENT DINNER HISTORY — avoid repeating meals from the past few weeks:
{history or "(no history yet)"}

MEAL PLAN FORMAT:
- Present the plan as a markdown table with columns: Day | Breakfast | Snack 1 | Lunch | Snack 2 | Dinner | Recipe Link
- Weekday lunch = leftovers from the previous night's dinner
- Leave Snack 1 and Snack 2 empty on weekends
- Follow all family rules strictly (nut-free weekday snacks/lunches/dinners, child preferences, etc.)
- Be conversational and easy to iterate with — keep responses concise

GENERATING EXCEL:
When the user says to generate the Excel, generate the file, or is happy with the plan, include this JSON block at the end of your response (it will be hidden from display):

```json
{{"plan": [
  {{"day": "Mon Jun 2", "breakfast": "...", "snack1": "...", "lunch": "...", "snack2": "...", "dinner": "...", "recipe_url": "..."}},
  ... (include all days of the week)
]}}
```"""


def clean_for_display(text: str) -> str:
    """Strip JSON plan blocks from text before showing to user."""
    # Remove complete blocks first
    text = re.sub(r'```json\s*\{.*?\}\s*```', '', text, flags=re.DOTALL)
    # Hide partial block still being streamed
    partial = text.find('```json')
    if partial != -1:
        text = text[:partial]
    return text.strip()


def extract_plan(text: str):
    """Extract the meal plan JSON from Claude's response, or return None."""
    match = re.search(r'```json\s*(\{.*?\})\s*```', text, re.DOTALL)
    if match:
        try:
            return json.loads(match.group(1)).get("plan")
        except Exception:
            return None
    return None


def call_claude(messages: list, notes: str, history: str, stream_placeholder):
    """Call Claude with streaming. Updates stream_placeholder in real-time."""
    client = anthropic.Anthropic(api_key=get_api_key())
    chunks = []
    with client.messages.stream(
        model="claude-opus-4-7",
        max_tokens=4096,
        system=[
            {
                "type": "text",
                "text": build_system_prompt(notes, history),
                # Prompt caching: large system prompt cached after first request
                "cache_control": {"type": "ephemeral"},
            }
        ],
        messages=messages,
    ) as stream:
        for text in stream.text_stream:
            chunks.append(text)
            full = "".join(chunks)
            stream_placeholder.markdown(clean_for_display(full) + "▌")
    return "".join(chunks)


def make_excel(plan: list) -> BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    HEADERS    = ["Day", "Breakfast", "Snack 1", "Lunch", "Snack 2", "Dinner", "Recipe Link"]
    COL_WIDTHS = [16, 32, 28, 36, 28, 60, 55]
    HEADER_FILL = PatternFill("solid", fgColor="4F81BD")
    HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
    ALT_FILL    = PatternFill("solid", fgColor="DCE6F1")
    PLAIN_FILL  = PatternFill("solid", fgColor="FFFFFF")
    WRAP        = Alignment(wrap_text=True, vertical="top")
    THIN        = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"),  bottom=Side(style="thin"),
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "Meal Plan"
    for col, (header, width) in enumerate(zip(HEADERS, COL_WIDTHS), 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment, cell.border = WRAP, THIN
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 20
    def extract_url(text):
        """Pull a bare URL out of 'Name | URL', '[name](url)', or a raw URL."""
        m = re.search(r'\[.*?\]\((https?://[^\)]+)\)', text)
        if m:
            return m.group(1)
        if " | http" in text:
            return text.split(" | ", 1)[1].strip()
        if text.strip().startswith("http"):
            return text.strip()
        return ""

    def parse_link(text):
        if " | http" in text:
            name, url = text.split(" | ", 1)
            return name.strip(), url.strip()
        m = re.search(r'\[(.+?)\]\((https?://[^\)]+)\)', text)
        if m:
            return m.group(1), m.group(2)
        return text, None

    def apply_link(cell, url):
        cell.hyperlink = url
        cell.font = Font(color="0563C1", underline="single")

    for i, day in enumerate(plan):
        row  = i + 2
        fill = ALT_FILL if i % 2 == 0 else PLAIN_FILL
        breakfast_text, breakfast_url = parse_link(day.get("breakfast", ""))
        snack1_text,    snack1_url    = parse_link(day.get("snack1", ""))
        snack2_text,    snack2_url    = parse_link(day.get("snack2", ""))
        display = {
            "day": day.get("day", ""),
            "breakfast": breakfast_text,
            "snack1": snack1_text,
            "lunch": day.get("lunch", ""),
            "snack2": snack2_text,
            "dinner": day.get("dinner", ""),
            "recipe_url": day.get("recipe_url", ""),
        }
        for col, key in enumerate(["day", "breakfast", "snack1", "lunch", "snack2", "dinner", "recipe_url"], 1):
            cell = ws.cell(row=row, column=col, value=display[key])
            cell.fill, cell.alignment, cell.border = fill, WRAP, THIN
        if breakfast_url:
            apply_link(ws.cell(row=row, column=2), breakfast_url)
        if snack1_url:
            apply_link(ws.cell(row=row, column=3), snack1_url)
        if snack2_url:
            apply_link(ws.cell(row=row, column=5), snack2_url)
        url = extract_url(day.get("recipe_url", ""))
        if url:
            apply_link(ws.cell(row=row, column=7), url)
        ws.row_dimensions[row].height = 60
    ws.page_setup.orientation  = "landscape"
    ws.page_setup.fitToPage    = True
    ws.page_setup.fitToWidth   = 1
    ws.page_setup.fitToHeight  = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ── Streamlit App ──
st.set_page_config(page_title="Family Meal Planner", page_icon="🥗", layout="wide")

def check_password():
    if st.session_state.get("authenticated"):
        return True
    pwd = st.text_input("Password", type="password", key="pwd_input")
    if st.button("Enter"):
        correct = os.getenv("APP_PASSWORD") or st.secrets.get("APP_PASSWORD", "")
        if pwd == correct:
            st.session_state.authenticated = True
            st.rerun()
        else:
            st.error("Incorrect password.")
    return False

if not check_password():
    st.stop()

if "notes" not in st.session_state:
    st.session_state.notes = load_text(
        APP_DIR / "meal-plan-notes.md",
        PROJECT_DIR / "meal-plan-notes.md",
    )
    st.session_state.history = load_text(
        APP_DIR / "meal_plan_history.md",
        MEMORY_DIR / "meal_plan_history.md",
        PROJECT_DIR / "meal_plan_history.md",
    )

notes, history = st.session_state.notes, st.session_state.history

for key, default in [("messages", []), ("plan_data", None), ("week_label", ""), ("history_saved", False)]:
    if key not in st.session_state:
        st.session_state[key] = default

with st.sidebar:
    st.title("🥗 Meal Planner")
    st.caption("Powered by Claude AI")
    st.divider()
    if get_api_key():
        st.success("API key loaded ✓", icon="🔑")
    else:
        st.error("Set ANTHROPIC_API_KEY in your .env file")
    st.divider()
    st.subheader("Plan a new week")
    today         = date.today()
    days_ahead    = (7 - today.weekday()) % 7 or 7
    default_start = today + timedelta(days=days_ahead)
    default_end   = default_start + timedelta(days=6)
    week_start = st.date_input("Week start (Mon)", value=default_start)
    week_end   = st.date_input("Week end (Sun)",   value=default_end)
    eating_out  = st.text_input("🍽️ Eating out?", placeholder="e.g. Friday dinner")
    busy_nights = st.text_input("⚡ Busy nights?", placeholder="e.g. Tuesday, Thursday")
    use_up      = st.text_input("🥦 Ingredients to use up?", placeholder="e.g. spinach, chicken")
    extra       = st.text_area("📝 Other notes", placeholder="Anything else Claude should know", height=80)
    if st.button("✨ Generate Plan", type="primary", use_container_width=True):
        parts = []
        if eating_out:  parts.append(f"eating out: {eating_out}")
        if busy_nights: parts.append(f"busy nights (need quick meals): {busy_nights}")
        if use_up:      parts.append(f"use up: {use_up}")
        if extra:       parts.append(extra)
        week_str = f"{week_start.strftime('%b %d')}–{week_end.strftime('%b %d, %Y')}"
        msg = f"Please create a meal plan for the week of {week_str}."
        if parts:
            msg += f" Constraints: {'; '.join(parts)}."
        st.session_state.messages      = []
        st.session_state.plan_data     = None
        st.session_state.history_saved = False
        st.session_state.week_label    = f"{week_start.strftime('%Y-%m-%d')}_to_{week_end.strftime('%m-%d')}"
        st.session_state.week_start    = week_start
        st.session_state.week_end      = week_end
        st.session_state.pending       = msg

st.title("Family Meal Planner")
if not st.session_state.messages and not st.session_state.get("pending"):
    st.info("Fill in the details on the left and click **✨ Generate Plan** to start.", icon="👈")

for msg in st.session_state.messages:
    with st.chat_message(msg["role"]):
        st.markdown(clean_for_display(msg["content"]))

pending    = st.session_state.pop("pending", None)
user_input = pending or st.chat_input("Ask for changes, or say 'generate Excel' when ready…")

if user_input:
    if not get_api_key():
        st.error("Please set ANTHROPIC_API_KEY in your .env file.")
        st.stop()
    st.session_state.messages.append({"role": "user", "content": user_input})
    with st.chat_message("user"):
        st.markdown(user_input)
    with st.chat_message("assistant"):
        placeholder = st.empty()
        reply = call_claude(
            messages=[{"role": m["role"], "content": m["content"]} for m in st.session_state.messages],
            notes=notes,
            history=history,
            stream_placeholder=placeholder,
        )
        placeholder.markdown(clean_for_display(reply))
        plan = extract_plan(reply)
        if plan:
            st.session_state.plan_data = plan
            st.success("Excel ready! Download below, then save to history when you're happy with it.", icon="✅")
    st.session_state.messages.append({"role": "assistant", "content": reply})
    st.rerun()

if st.session_state.plan_data:
    st.divider()
    col1, col2, col3 = st.columns([2, 1, 1])
    with col1:
        st.subheader("📊 Your meal plan is ready to download")
    with col2:
        filename = f"{st.session_state.week_label or 'meal-plan'}.xlsx"
        buf = make_excel(st.session_state.plan_data)
        st.download_button(
            label="⬇️ Download Excel",
            data=buf,
            file_name=filename,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
            use_container_width=True,
        )
    with col3:
        if st.session_state.history_saved:
            st.success("Saved to history ✓")
        elif get_github_token():
            if st.button("💾 Save to History", use_container_width=True):
                entry = build_history_entry(
                    st.session_state.plan_data,
                    st.session_state.get("week_start", date.today()),
                    st.session_state.get("week_end", date.today()),
                )
                ok, msg = save_history_to_github(entry)
                if ok:
                    st.session_state.history_saved = True
                    st.rerun()
                else:
                    st.error(msg)
        else:
            st.caption("Add GITHUB_TOKEN to secrets to enable auto-save.")
